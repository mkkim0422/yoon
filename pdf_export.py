"""
pdf_export.py — Excel Invoice 시트를 PDF로 변환.

구현 방식 — OS 분기:
  · Windows: Excel COM (pywin32). Streamlit 워커 스레드의 COM apartment 문제를
    피하려고 별도 Python subprocess 에서 Dispatch.
  · Linux  : LibreOffice headless (`soffice --convert-to pdf`). Streamlit Cloud
    같은 Linux 환경 대응. 변환 전에 openpyxl 로 sheet_name 외 시트는 hidden
    처리하고 A4·여백·"한 페이지에 맞춤" 설정을 주입해 Excel COM 과 가능한
    가까운 레이아웃을 얻는다.

요구 사항:
  · Windows: Microsoft Excel + pywin32
  · Linux  : libreoffice (apt: `libreoffice-calc`, `libreoffice-core`)
"""
from __future__ import annotations

import os
import platform
import shutil
import subprocess
import sys
import tempfile
import uuid
from pathlib import Path


# subprocess 안에서 실행될 Excel COM 스크립트.
# xlsx → PDF 변환만 수행; 결과/오류는 stdout/stderr 로 보고.
_CHILD_SCRIPT = r'''
import sys, os, traceback
try:
    import pythoncom
    import win32com.client as win32
except Exception as e:
    sys.stderr.write(f"pywin32 import failed: {e}\n")
    sys.exit(2)

XLSX_PATH  = sys.argv[1]
PDF_PATH   = sys.argv[2]
SHEET_NAME = sys.argv[3]

XL_TYPE_PDF  = 0
XL_PORTRAIT  = 1
XL_LANDSCAPE = 2
XL_STANDARD  = 0
XL_PAPER_A4  = 9

excel = None
wb    = None
try:
    pythoncom.CoInitialize()
    try:
        excel = win32.DispatchEx("Excel.Application")
    except Exception as e:
        sys.stderr.write(f"Excel dispatch failed: {e}\n")
        sys.exit(3)

    excel.Visible       = False
    excel.DisplayAlerts = False
    try:
        excel.AskToUpdateLinks   = False
        excel.AutomationSecurity = 3
    except Exception:
        pass

    wb = excel.Workbooks.Open(os.path.abspath(XLSX_PATH),
                              ReadOnly=True, UpdateLinks=0)
    try:
        excel.CalculateFull()
    except Exception:
        pass

    try:
        ws = wb.Worksheets(SHEET_NAME)
    except Exception as e:
        sys.stderr.write(f"Worksheet '{SHEET_NAME}' not found: {e}\n")
        sys.exit(4)

    try:
        # 세로 A4, 가로/세로 모두 1페이지에 맞춤 (첨부 샘플 PDF 동일 레이아웃)
        ws.PageSetup.PaperSize          = XL_PAPER_A4
        ws.PageSetup.Orientation        = XL_PORTRAIT
        ws.PageSetup.Zoom               = False
        ws.PageSetup.FitToPagesWide     = 1
        ws.PageSetup.FitToPagesTall     = 1     # 한 페이지 안에 높이도 맞춤
        ws.PageSetup.CenterHorizontally = True
        ws.PageSetup.CenterVertically   = False
        # 여백 최소화 — 빈 공간 감소
        ws.PageSetup.LeftMargin   = excel.InchesToPoints(0.25)
        ws.PageSetup.RightMargin  = excel.InchesToPoints(0.25)
        ws.PageSetup.TopMargin    = excel.InchesToPoints(0.3)
        ws.PageSetup.BottomMargin = excel.InchesToPoints(0.3)
        ws.PageSetup.HeaderMargin = excel.InchesToPoints(0.1)
        ws.PageSetup.FooterMargin = excel.InchesToPoints(0.1)
        # 프린트 영역을 실제 콘텐츠 범위로 좁히기 (빈 행으로 인한 추가 페이지 방지)
        try:
            used = ws.UsedRange
            ws.PageSetup.PrintArea = used.Address
        except Exception:
            pass
    except Exception:
        pass

    ws.Select()

    ws.ExportAsFixedFormat(
        Type=XL_TYPE_PDF,
        Filename=os.path.abspath(PDF_PATH),
        Quality=XL_STANDARD,
        IncludeDocProperties=True,
        IgnorePrintAreas=False,
        OpenAfterPublish=False,
    )

    if not os.path.exists(PDF_PATH) or os.path.getsize(PDF_PATH) == 0:
        sys.stderr.write("Excel returned but PDF was not written.\n")
        sys.exit(5)

    sys.stdout.write("OK\n")
    sys.exit(0)

except SystemExit:
    raise
except Exception:
    sys.stderr.write(traceback.format_exc())
    sys.exit(1)
finally:
    try:
        if wb is not None:
            wb.Close(SaveChanges=False)
    except Exception:
        pass
    try:
        if excel is not None:
            excel.Quit()
    except Exception:
        pass
    try:
        pythoncom.CoUninitialize()
    except Exception:
        pass
'''


def _find_soffice() -> str | None:
    """Linux/Mac 의 LibreOffice 바이너리 경로. PATH 우선, 없으면 흔한 위치 탐색."""
    for cand in ("soffice", "libreoffice"):
        p = shutil.which(cand)
        if p:
            return p
    for p in (
        "/usr/bin/soffice", "/usr/bin/libreoffice",
        "/usr/lib/libreoffice/program/soffice",
        "/snap/bin/libreoffice",
    ):
        if Path(p).exists():
            return p
    return None


def _xlsx_to_pdf_libreoffice(
    xlsx_bytes: bytes,
    sheet_name: str,
    timeout_sec: int,
) -> tuple[bytes | None, str | None]:
    """LibreOffice headless 로 xlsx → PDF. sheet_name 외 시트는 숨겨서 첫 페이지가
    해당 시트가 되도록 한다. PageSetup(A4·여백·1페이지 맞춤)도 openpyxl 로 주입.
    """
    soffice = _find_soffice()
    if not soffice:
        return None, (
            "LibreOffice(soffice) 를 찾을 수 없습니다. "
            "packages.txt 에 `libreoffice-calc` 가 포함되어 있는지 확인하세요."
        )

    base = Path(tempfile.gettempdir()) / f"sph_pdf_{uuid.uuid4().hex[:10]}"
    base.mkdir(parents=True, exist_ok=True)
    xlsx_path = base / "invoice.xlsx"
    # 출력 PDF 는 soffice 가 입력 파일명 기준 같은 이름으로 만든다(invoice.pdf)
    pdf_path = base / "invoice.pdf"

    try:
        xlsx_path.write_bytes(xlsx_bytes)

        # ── sheet_name 외 시트 숨김 + PageSetup 주입 ──────────────────
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_path)
            if sheet_name not in wb.sheetnames:
                return None, f"엑셀 파일에 '{sheet_name}' 시트가 없습니다."
            # 대상 시트를 active 로 설정 + 다른 시트는 hidden
            for nm in wb.sheetnames:
                ws = wb[nm]
                if nm == sheet_name:
                    wb.active = wb.sheetnames.index(nm)
                else:
                    ws.sheet_state = "hidden"
            ws = wb[sheet_name]
            # 페이지 설정 — Excel COM 분기와 동일한 의도(A4, 1페이지 fit, 좁은 여백)
            ws.page_setup.orientation     = ws.ORIENTATION_PORTRAIT
            ws.page_setup.paperSize       = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth      = 1
            ws.page_setup.fitToHeight     = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_options.horizontalCentered = True
            ws.page_margins.left   = 0.25
            ws.page_margins.right  = 0.25
            ws.page_margins.top    = 0.3
            ws.page_margins.bottom = 0.3
            ws.page_margins.header = 0.1
            ws.page_margins.footer = 0.1
            wb.save(xlsx_path)
        except Exception as e:
            # openpyxl 처리 실패해도 그대로 변환 시도 (품질만 낮아질 뿐).
            sys.stderr.write(f"pre-pdf openpyxl tweak failed: {e}\n")

        # ── LibreOffice headless 변환 ───────────────────────────────
        # HOME 환경변수 없는 컨테이너 환경 대응 — --env:UserInstallation
        user_profile = base / "lo_profile"
        completed = subprocess.run(
            [
                soffice,
                "--headless",
                "--norestore", "--nofirststartwizard", "--nologo",
                f"-env:UserInstallation=file://{user_profile}",
                "--convert-to", "pdf",
                "--outdir", str(base),
                str(xlsx_path),
            ],
            capture_output=True,
            timeout=timeout_sec,
        )

        if pdf_path.exists() and pdf_path.stat().st_size > 0:
            return pdf_path.read_bytes(), None

        stderr = (completed.stderr or b"").decode(errors="ignore").strip()
        stdout = (completed.stdout or b"").decode(errors="ignore").strip()
        return None, (
            f"LibreOffice 변환 실패 [rc={completed.returncode}]: "
            f"{stderr or stdout or '(추가 정보 없음)'}"
        )
    except subprocess.TimeoutExpired:
        return None, f"PDF 변환이 {timeout_sec} 초 내에 완료되지 않았습니다."
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"
    finally:
        for p in (xlsx_path, pdf_path):
            try:
                if p.exists(): p.unlink()
            except Exception:
                pass
        try:
            shutil.rmtree(base, ignore_errors=True)
        except Exception:
            pass


def xlsx_sheet_to_pdf(
    xlsx_bytes: bytes,
    sheet_name: str = "Invoice",
    timeout_sec: int = 90,
) -> tuple[bytes | None, str | None]:
    """xlsx 바이트에서 지정 시트를 PDF 바이트로 추출.

    반환:
      (pdf_bytes, None)  — 성공
      (None, err_msg)    — 실패 (err_msg 는 사용자에게 표시할 원인)
    """
    # Linux/macOS → LibreOffice 분기 (Streamlit Cloud 등 Excel 없는 환경 대응).
    if platform.system() != "Windows":
        # LibreOffice 가 무거운 변환이라 기본 timeout 을 좀 더 넉넉히.
        return _xlsx_to_pdf_libreoffice(xlsx_bytes, sheet_name, max(timeout_sec, 120))

    # 임시 작업 디렉터리 — Excel COM 은 한글/공백 경로에 약하므로 %TEMP% 하위 영문 경로 사용
    base = Path(tempfile.gettempdir()) / f"sph_pdf_{uuid.uuid4().hex[:10]}"
    base.mkdir(parents=True, exist_ok=True)
    xlsx_path = base / "invoice.xlsx"
    pdf_path  = base / "invoice.pdf"
    script_path = base / "_conv.py"

    pdf_bytes: bytes | None = None
    err_msg:   str   | None = None

    try:
        xlsx_path.write_bytes(xlsx_bytes)
        script_path.write_text(_CHILD_SCRIPT, encoding="utf-8")

        # subprocess 로 격리 실행 (Streamlit 의 COM apartment 와 완전 분리)
        completed = subprocess.run(
            [sys.executable, str(script_path),
             str(xlsx_path), str(pdf_path), sheet_name],
            capture_output=True,
            timeout=timeout_sec,
            # Windows 에서 콘솔 창이 뜨지 않도록
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )

        if completed.returncode == 0 and pdf_path.exists() and pdf_path.stat().st_size > 0:
            pdf_bytes = pdf_path.read_bytes()
        else:
            stderr = (completed.stderr or b"").decode(errors="ignore").strip()
            stdout = (completed.stdout or b"").decode(errors="ignore").strip()
            rc = completed.returncode
            hint = {
                2: "pywin32 가 설치되지 않았거나 COM 등록이 필요합니다 "
                   "(`python -m pip install --upgrade pywin32` 후 "
                   "`python Scripts/pywin32_postinstall.py -install`).",
                3: "Microsoft Excel 이 설치되지 않았거나 현재 계정에서 실행 불가합니다.",
                4: f"엑셀 파일에 '{sheet_name}' 시트가 존재하지 않습니다.",
                5: "Excel 이 PDF 를 생성하지 못했습니다 (프린터 드라이버 확인).",
            }.get(rc, "원인 불명")
            err_msg = f"[rc={rc}] {hint}\n{stderr or stdout or '(추가 정보 없음)'}"

    except subprocess.TimeoutExpired:
        err_msg = f"PDF 변환이 {timeout_sec} 초 내에 완료되지 않았습니다."
    except FileNotFoundError as e:
        err_msg = f"Python 실행 파일을 찾지 못했습니다: {e}"
    except Exception as e:
        err_msg = f"{type(e).__name__}: {e}"
    finally:
        # 임시 파일 정리
        for p in (xlsx_path, pdf_path, script_path):
            try:
                if p.exists(): p.unlink()
            except Exception:
                pass
        try:
            os.rmdir(base)
        except Exception:
            pass

    return pdf_bytes, err_msg


# ── 일괄 PDF 변환: Excel 인스턴스 1개를 batch 동안 살려두고 재사용 ──────────
# 단일 호출 xlsx_sheet_to_pdf() 는 회사마다 subprocess + Excel 시작/종료가
# 일어나 회사당 11~40초가 소요됨. BatchExcelPdf 컨텍스트 매니저는 subprocess
# 와 Excel.Application 을 batch 동안 1회만 띄워두고 stdin 으로 변환 명령을
# 흘려보내, 2번째 호출부터 회사당 2~4초로 단축한다.
#
# 사용:
#     with BatchExcelPdf() as conv:
#         for company in companies:
#             pdf_bytes, err = conv.convert(xlsx_bytes, sheet_name)
#
# Linux 환경: 서버 모드 미구현 — fallback 으로 _xlsx_to_pdf_libreoffice 호출.
# subprocess/Excel 시작 실패 시에도 단일 호출로 graceful fallback.

_BATCH_SERVER_SCRIPT = r'''
"""Excel COM long-lived 서버. stdin 으로 \t 구분 명령 수신, stdout 으로 결과 송신.

프로토콜:
  IN:  <xlsx_path>\t<pdf_path>\t<sheet_name>\n
  OUT: OK\n                — 성공
       ERR\t<message>\n   — 실패
  EXIT: 빈 줄 또는 "EXIT" → 정상 종료
"""
import sys, os, traceback

try:
    import pythoncom
    import win32com.client as win32
except Exception as e:
    sys.stderr.write(f"pywin32 import failed: {e}\n")
    sys.exit(2)

XL_TYPE_PDF = 0
XL_PORTRAIT = 1
XL_STANDARD = 0
XL_PAPER_A4 = 9

excel = None
try:
    pythoncom.CoInitialize()
    try:
        excel = win32.DispatchEx("Excel.Application")
    except Exception as e:
        sys.stderr.write(f"Excel dispatch failed: {e}\n")
        sys.exit(3)

    excel.Visible       = False
    excel.DisplayAlerts = False
    try:
        excel.AskToUpdateLinks   = False
        excel.AutomationSecurity = 3
    except Exception:
        pass

    sys.stdout.write("READY\n")
    sys.stdout.flush()

    while True:
        line = sys.stdin.readline()
        if not line:
            break
        line = line.rstrip("\r\n")
        if not line or line == "EXIT":
            break
        parts = line.split("\t")
        if len(parts) != 3:
            sys.stdout.write(f"ERR\tinvalid command\n")
            sys.stdout.flush()
            continue
        xlsx_path, pdf_path, sheet_name = parts
        wb = None
        try:
            wb = excel.Workbooks.Open(
                os.path.abspath(xlsx_path), ReadOnly=True, UpdateLinks=0
            )
            try:
                excel.CalculateFull()
            except Exception:
                pass
            try:
                ws = wb.Worksheets(sheet_name)
            except Exception as e:
                sys.stdout.write(f"ERR\tWorksheet '{sheet_name}' not found: {e}\n")
                sys.stdout.flush()
                continue

            try:
                ws.PageSetup.PaperSize          = XL_PAPER_A4
                ws.PageSetup.Orientation        = XL_PORTRAIT
                ws.PageSetup.Zoom               = False
                ws.PageSetup.FitToPagesWide     = 1
                ws.PageSetup.FitToPagesTall     = 1
                ws.PageSetup.CenterHorizontally = True
                ws.PageSetup.CenterVertically   = False
                ws.PageSetup.LeftMargin   = excel.InchesToPoints(0.25)
                ws.PageSetup.RightMargin  = excel.InchesToPoints(0.25)
                ws.PageSetup.TopMargin    = excel.InchesToPoints(0.3)
                ws.PageSetup.BottomMargin = excel.InchesToPoints(0.3)
                ws.PageSetup.HeaderMargin = excel.InchesToPoints(0.1)
                ws.PageSetup.FooterMargin = excel.InchesToPoints(0.1)
                try:
                    used = ws.UsedRange
                    ws.PageSetup.PrintArea = used.Address
                except Exception:
                    pass
            except Exception:
                pass

            ws.Select()
            ws.ExportAsFixedFormat(
                Type=XL_TYPE_PDF,
                Filename=os.path.abspath(pdf_path),
                Quality=XL_STANDARD,
                IncludeDocProperties=True,
                IgnorePrintAreas=False,
                OpenAfterPublish=False,
            )

            if not os.path.exists(pdf_path) or os.path.getsize(pdf_path) == 0:
                sys.stdout.write("ERR\tExcel returned but PDF was not written\n")
                sys.stdout.flush()
            else:
                sys.stdout.write("OK\n")
                sys.stdout.flush()
        except Exception as e:
            err_short = f"{type(e).__name__}: {e}".replace("\t", " ").replace("\n", " ")
            sys.stdout.write(f"ERR\t{err_short}\n")
            sys.stdout.flush()
        finally:
            try:
                if wb is not None:
                    wb.Close(SaveChanges=False)
            except Exception:
                pass

except Exception:
    sys.stderr.write(traceback.format_exc())
finally:
    try:
        if excel is not None:
            excel.Quit()
    except Exception:
        pass
    try:
        pythoncom.CoUninitialize()
    except Exception:
        pass
'''


def _list_excel_pids() -> set[int]:
    """현재 떠 있는 EXCEL.EXE PID 집합. Windows 전용. 실패 시 빈 집합."""
    if platform.system() != "Windows":
        return set()
    try:
        out = subprocess.check_output(
            ["tasklist", "/FI", "IMAGENAME eq EXCEL.EXE", "/FO", "CSV", "/NH"],
            text=True, stderr=subprocess.DEVNULL, encoding="utf-8", errors="replace",
            timeout=5,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        pids: set[int] = set()
        for line in out.splitlines():
            line = line.strip()
            if not line or "EXCEL.EXE" not in line.upper():
                continue
            parts = line.split(",")
            if len(parts) >= 2:
                pid_str = parts[1].strip().strip('"')
                try:
                    pids.add(int(pid_str))
                except ValueError:
                    pass
        return pids
    except Exception:
        return set()


def _kill_pids(pids: set[int]) -> int:
    """주어진 PID 들을 taskkill /F 로 강제 종료. 성공 카운트 반환."""
    if not pids or platform.system() != "Windows":
        return 0
    killed = 0
    for pid in pids:
        try:
            r = subprocess.run(
                ["taskkill", "/F", "/PID", str(pid)],
                capture_output=True, timeout=5,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            if r.returncode == 0:
                killed += 1
        except Exception:
            pass
    return killed


class BatchExcelPdf:
    """Windows 전용. with 블록 동안 Excel 1개를 살려두고 PDF 변환 반복.

    Linux/macOS 에서는 서버 모드 미구현. convert() 가 자동으로
    _xlsx_to_pdf_libreoffice() 로 fallback. Excel COM/subprocess 시작 실패
    시에도 단일 호출 xlsx_sheet_to_pdf() 로 graceful fallback.

    Excel COM 좀비 누수 방지: __enter__ 직전 EXCEL.EXE PID 스냅샷을 떠두고,
    __exit__ 시 새로 생긴 PID 만 강제 종료. 정상 종료 경로에서 excel.Quit()
    이 제대로 안 끝나는 경우(컴퓨터가 바빠 timeout, COM 참조 잔여 등)에도
    좀비를 남기지 않음. 배치 반복 시 누적 누수로 시스템 슬로우다운 방지.
    """

    def __init__(self, timeout_per_call: int = 120):
        self.timeout = timeout_per_call
        self._proc = None
        self._base: Path | None = None
        self._is_windows = (platform.system() == "Windows")
        self._call_idx = 0
        self._excel_pids_before: set[int] = set()
        self._our_excel_pids: set[int] = set()

    def __enter__(self):
        if self._is_windows:
            self._start_windows_server()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self._stop_windows_server()
        return False

    def _start_windows_server(self) -> None:
        try:
            # Excel PID 스냅샷 — 우리가 띄운 Excel 만 추적하기 위함
            self._excel_pids_before = _list_excel_pids()

            self._base = Path(tempfile.gettempdir()) / f"sph_pdf_batch_{uuid.uuid4().hex[:10]}"
            self._base.mkdir(parents=True, exist_ok=True)
            script_path = self._base / "_server.py"
            script_path.write_text(_BATCH_SERVER_SCRIPT, encoding="utf-8")

            self._proc = subprocess.Popen(
                [sys.executable, "-u", str(script_path)],
                stdin=subprocess.PIPE,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                bufsize=1,
                text=True,
                encoding="utf-8",
                errors="replace",
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
            # READY 대기 — Excel 시작이 느릴 수 있어 30초까지 허용
            ready = self._read_line_with_timeout(30)
            if ready != "READY":
                # 실패 — 서버 정리
                self._stop_windows_server()
                return
            # READY 받았으면 새로 생긴 Excel PID 들이 우리 것
            after = _list_excel_pids()
            self._our_excel_pids = after - self._excel_pids_before
        except Exception as e:
            sys.stderr.write(f"BatchExcelPdf server start failed: {e}\n")
            self._stop_windows_server()

    def _stop_windows_server(self) -> None:
        if self._proc is not None:
            try:
                if self._proc.poll() is None and self._proc.stdin and not self._proc.stdin.closed:
                    try:
                        self._proc.stdin.write("EXIT\n")
                        self._proc.stdin.flush()
                    except Exception:
                        pass
                try:
                    self._proc.wait(timeout=10)
                except Exception:
                    try: self._proc.kill()
                    except Exception: pass
            except Exception:
                pass
            self._proc = None
        # Excel COM 좀비 강제 정리 — Quit() 이 행되거나 subprocess 강제 종료
        # 된 경우에도 우리가 띄운 Excel 만 골라서 taskkill. 사용자가 별도로
        # 띄운 Excel 은 _excel_pids_before 에 포함되어 있어 건드리지 않음.
        if self._our_excel_pids:
            alive = self._our_excel_pids & _list_excel_pids()
            if alive:
                killed = _kill_pids(alive)
                if killed > 0:
                    print(f"[BatchExcelPdf] 좀비 Excel {killed}개 강제 정리 (PID={sorted(alive)})", flush=True)
            self._our_excel_pids = set()
        if self._base is not None:
            try:
                shutil.rmtree(self._base, ignore_errors=True)
            except Exception:
                pass
            self._base = None

    def _read_line_with_timeout(self, timeout_sec: float):
        """Windows pipe 에서 timeout 부 readline. 스레드+queue 패턴."""
        if self._proc is None or self._proc.stdout is None:
            return None
        import threading, queue as _q
        result_q: _q.Queue = _q.Queue()
        def _reader():
            try:
                line = self._proc.stdout.readline()
                result_q.put(line)
            except Exception:
                result_q.put(None)
        t = threading.Thread(target=_reader, daemon=True)
        t.start()
        try:
            line = result_q.get(timeout=timeout_sec)
            return None if line is None else line.rstrip("\r\n")
        except _q.Empty:
            return None

    def convert(self, xlsx_bytes: bytes, sheet_name: str = "Invoice") -> tuple[bytes | None, str | None]:
        """xlsx_bytes 를 PDF 로 변환. 서버 살아있으면 재사용, 아니면 fallback."""
        # 서버 미가용 → 단일 호출로 fallback
        if (not self._is_windows) or self._proc is None or self._proc.poll() is not None:
            return xlsx_sheet_to_pdf(xlsx_bytes, sheet_name, self.timeout)

        self._call_idx += 1
        cid = f"{self._call_idx:04d}_{uuid.uuid4().hex[:6]}"
        assert self._base is not None
        xlsx_path = self._base / f"in_{cid}.xlsx"
        pdf_path  = self._base / f"out_{cid}.pdf"

        try:
            xlsx_path.write_bytes(xlsx_bytes)
            cmd = f"{xlsx_path}\t{pdf_path}\t{sheet_name}\n"
            self._proc.stdin.write(cmd)
            self._proc.stdin.flush()

            resp = self._read_line_with_timeout(self.timeout)
            if resp is None:
                # 타임아웃 — 서버 죽었을 가능성 큼. 정리 후 단일 호출 fallback.
                self._stop_windows_server()
                return xlsx_sheet_to_pdf(xlsx_bytes, sheet_name, self.timeout)

            if resp.startswith("OK"):
                if pdf_path.exists() and pdf_path.stat().st_size > 0:
                    pdf_bytes = pdf_path.read_bytes()
                    return pdf_bytes, None
                return None, "Excel returned OK but PDF not written"
            elif resp.startswith("ERR"):
                err = resp.split("\t", 1)[1] if "\t" in resp else "unknown"
                return None, f"[batch] {err}"
            else:
                return None, f"[batch] unexpected response: {resp[:200]}"
        except Exception as e:
            # 통신 실패 — 서버 정리 후 fallback
            self._stop_windows_server()
            return xlsx_sheet_to_pdf(xlsx_bytes, sheet_name, self.timeout)
        finally:
            for p in (xlsx_path, pdf_path):
                try:
                    if p.exists(): p.unlink()
                except Exception:
                    pass


def is_available() -> bool:
    """PDF 변환 가능 여부.
       Windows: subprocess 로 Excel dispatch 시도.
       Linux/macOS: LibreOffice(soffice) 바이너리 존재 확인.
    """
    if platform.system() != "Windows":
        return _find_soffice() is not None

    script = (
        "import sys\n"
        "try:\n"
        "    import pythoncom, win32com.client as w\n"
        "    pythoncom.CoInitialize()\n"
        "    e = w.DispatchEx('Excel.Application')\n"
        "    e.Quit()\n"
        "    print('OK')\n"
        "except Exception as x:\n"
        "    sys.stderr.write(str(x))\n"
        "    sys.exit(1)\n"
    )
    try:
        r = subprocess.run(
            [sys.executable, "-c", script],
            capture_output=True, timeout=15,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        return r.returncode == 0
    except Exception:
        return False
