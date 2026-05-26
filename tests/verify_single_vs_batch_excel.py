"""파일 A(개별 정산) vs 파일 B(일괄 정산) 엑셀 정밀 비교 검증.

생성 경로:
  A) 단일 정산 — preprocess → load → engine → generate_formatted_invoice
     를 직접 순차 호출 (webapp 단일 모드 코드와 동일 함수 체인).
  B) 일괄 정산 — webapp._run_batch_single_billing 단일 호출 (zip 안 파일과 동일).

비교 단계:
  1) 두 엑셀의 시트 수·이름·각 시트의 컬럼 폭/행 수 출력 + 구조 일치 확인
  2) 시트별 셀 좌표(행, 열) 기준으로 모든 셀 값 비교
     - 숫자: round(., 6) 으로 부동소수점 잡음 흡수, 정수와 float 의 dtype 차이 무시
     - 문자열: 원본 그대로 비교 (공백/줄바꿈 차이도 포착)
     - None 셀은 빈 문자열과 동일 취급 (엑셀 디폴트)
  3) 결과 리포트:
     - 일치 시: "전체 N개사, 모든 셀 100% 일치"
     - 불일치 시: 회사별로 (시트, 좌표, 값A, 값B) 목록 + 불일치 건수 총합
"""
from __future__ import annotations

import io
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

# ── Streamlit 무력화 (webapp import 부작용 차단) ──────────────────────
class _AttrDict(dict):
    def __getattr__(self, k):
        if k in self:
            return self[k]
        raise AttributeError(k)
    def __setattr__(self, k, v):
        self[k] = v


class _StubST:
    def __init__(self):
        self.__dict__["session_state"] = _AttrDict()
    def __getattr__(self, k): return _StubST()
    def __call__(self, *a, **kw): return _StubST()
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def __contains__(self, k): return False
    def __setitem__(self, k, v): pass
    def __getitem__(self, k): return None
    def __iter__(self): return iter([])


sys.modules["streamlit"] = _StubST()
sys.modules["streamlit_sortables"] = _StubST()


def _extract_run_batch_single_billing():
    import ast
    src = (ROOT / "webapp.py").read_text(encoding="utf-8")
    tree = ast.parse(src)
    target = next(
        (n for n in tree.body
         if isinstance(n, ast.FunctionDef) and n.name == "_run_batch_single_billing"),
        None,
    )
    assert target, "_run_batch_single_billing not found"
    func_src = ast.get_source_segment(src, target)
    from pathlib import Path as _P
    from decimal import Decimal as _Dec
    from collections import defaultdict as _dd
    from billing.engine import calculate_billing as _cb, calculate_billing_by_project as _cbp
    from billing.loader import (
        build_sku_master_from_usage as _bsmf,
        detect_missing_skus as _dms,
        get_free_caps_from_price_list as _gfc,
        load_usage_rows as _lur,
    )
    from billing.models import BillingLineItem as _BLI
    from billing.preprocessor import preprocess_usage_file as _puf
    from invoice_generator import generate_formatted_invoice as _gfi
    ns = {
        "Path": _P, "Decimal": _Dec,
        "calculate_billing": _cb, "calculate_billing_by_project": _cbp,
        "build_sku_master_from_usage": _bsmf, "detect_missing_skus": _dms,
        "get_free_caps_from_price_list": _gfc, "load_usage_rows": _lur,
        "preprocess_usage_file": _puf, "generate_formatted_invoice": _gfi,
    }
    exec(func_src, ns)
    return ns["_run_batch_single_billing"]


_run_batch_single_billing = _extract_run_batch_single_billing()

from billing.engine import calculate_billing, calculate_billing_by_project  # noqa: E402
from billing.loader import (  # noqa: E402
    build_sku_master_from_usage, get_billable_sku_names, load_usage_rows,
)
from billing.preprocessor import extract_company_names, preprocess_usage_file  # noqa: E402
from invoice_generator import generate_formatted_invoice  # noqa: E402


# ── 1) 두 경로로 같은 회사의 엑셀 생성 ─────────────────────────────────
def gen_single(company, billing_month, csv_path, plist, currency,
               rate, bank, phrase, rate_date):
    raw = preprocess_usage_file(str(csv_path), billing_month, company_filter=company)
    usage = load_usage_rows(raw)
    master = build_sku_master_from_usage(usage, str(plist))
    billable = get_billable_sku_names(str(plist))
    ex = Decimal(str(rate))
    mr = Decimal("0")
    items = calculate_billing(usage, master, ex, mr, mode="account")
    proj = calculate_billing_by_project(usage, master, ex, mr, mode="account")
    return generate_formatted_invoice(
        line_items=items, company_name=company, billing_month=billing_month,
        exchange_rate=ex, margin_rate=mr, bank_name=bank, proj_results=proj,
        price_list_file=str(plist), sku_order=None, currency=currency,
        billable_skus=billable, billing_mode="account",
        per_project_invoices=None, min_charge_amount=0.0,
        min_charge_currency="KRW", rate_date_str=rate_date,
        rate_phrase=phrase, rate_extra="", include_project_sheet=True,
        subtotal_round=0 if currency == "KRW" else 2, force_keep_skus=None,
    )


def gen_batch(company, billing_month, csv_path, plist, currency,
              rate, bank, phrase, rate_date):
    billable = get_billable_sku_names(str(plist))
    res = _run_batch_single_billing(
        selected_company=company, billing_month=billing_month,
        tmp_input_path=Path(csv_path), price_list_file=str(plist),
        currency=currency, exchange_rate=float(rate), margin_rate=0.0,
        rate_date_str=rate_date, billing_mode="account",
        include_project_sheet=True,
        subtotal_round=0 if currency == "KRW" else 2,
        bank_name=bank, rate_phrase_text=phrase, rate_extra_text="",
        min_charge_amount=0.0, min_charge_currency="KRW",
        sku_order=[], manual_skus=[], hidden_skus=[],
        billable_skus=billable, dl_xlsx=True, dl_pdf=False,
    )
    assert res["ok"], f"batch failed: {res.get('error')}"
    return res["excel_bytes"]


# ── 2) 엑셀 구조 + 셀 단위 비교 ─────────────────────────────────────────
def normalize_cell(v):
    """비교용 셀 정규화 — float 잡음 흡수, None=빈문자열."""
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        # 정수와 동일한 float 은 정수로
        try:
            f = float(v)
            if f == int(f):
                return int(f)
            return round(f, 6)
        except (ValueError, TypeError):
            return v
    return v  # 문자열은 원본 그대로 (공백/줄바꿈 보존)


def load_excel_structure(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    info = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        info[sn] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "cells": {},  # (row, col) -> value
        }
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                info[sn]["cells"][(r, c)] = normalize_cell(v)
    return info


def compare_excels(a_bytes, b_bytes):
    a = load_excel_structure(a_bytes)
    b = load_excel_structure(b_bytes)
    # 구조 비교
    a_sheets = list(a.keys())
    b_sheets = list(b.keys())
    structure_match = (a_sheets == b_sheets)
    sheet_diffs = {}
    for sn in set(a_sheets) | set(b_sheets):
        if sn not in a:
            sheet_diffs[sn] = "missing in A"
        elif sn not in b:
            sheet_diffs[sn] = "missing in B"
        elif (a[sn]["max_row"], a[sn]["max_col"]) != (b[sn]["max_row"], b[sn]["max_col"]):
            sheet_diffs[sn] = (
                f"size A={a[sn]['max_row']}x{a[sn]['max_col']} "
                f"B={b[sn]['max_row']}x{b[sn]['max_col']}"
            )
    # 셀 비교
    cell_diffs = []
    for sn in a_sheets:
        if sn not in b:
            continue
        coords = set(a[sn]["cells"].keys()) | set(b[sn]["cells"].keys())
        for rc in sorted(coords):
            va = a[sn]["cells"].get(rc, "")
            vb = b[sn]["cells"].get(rc, "")
            if va != vb:
                cell_diffs.append((sn, rc, va, vb))
    return {
        "a_sheets": a_sheets, "b_sheets": b_sheets,
        "structure_match": structure_match,
        "sheet_diffs": sheet_diffs,
        "cell_diffs": cell_diffs,
        "a_info": {sn: (a[sn]["max_row"], a[sn]["max_col"]) for sn in a_sheets},
        "b_info": {sn: (b[sn]["max_row"], b[sn]["max_col"]) for sn in b_sheets},
    }


# ── 3) 메인: 다수 회사에 대해 비교 + 결과 리포트 ─────────────────────────
def main():
    csv = ROOT / "billing.csv"
    plist_usd = ROOT / "billing" / "saved_price_list_usd.xlsx"
    plist_krw = ROOT / "billing" / "saved_price_list_krw.xlsx"
    assert csv.exists(), f"missing {csv}"

    companies = extract_company_names(str(csv))

    # 정산 공통 파라미터
    billing_month = "2026-04"
    rate = 1480.80
    bank = "하나은행"
    phrase = "최종 송금환율 기준"
    rate_date = "2026.04.30"

    # 데이터 있는 회사 선별
    targets = []
    for c in companies[:30]:
        try:
            raw = preprocess_usage_file(str(csv), billing_month, company_filter=c)
            if raw:
                targets.append(c)
        except Exception:
            pass
        if len(targets) >= 12:
            break

    print(f"=== Comparing {len(targets)} companies ===\n")

    total_cell_diffs = 0
    only_in_a = []
    only_in_b = []
    company_results = []

    for c in targets:
        # 통화 자동 결정 — USD 우선 시도, 데이터 없으면 KRW
        used_currency = None
        used_plist = None
        for currency, plist in (("USD", plist_usd), ("KRW", plist_krw)):
            try:
                a_bytes = gen_single(
                    c, billing_month, csv, plist, currency,
                    rate, bank, phrase, rate_date,
                )
                # 빈 엑셀이면 다음 통화로
                wb = openpyxl.load_workbook(io.BytesIO(a_bytes), data_only=False)
                has_data = any(
                    wb[sn].max_row > 5 for sn in wb.sheetnames
                )
                if has_data:
                    used_currency = currency
                    used_plist = plist
                    break
            except Exception:
                continue
        if not used_currency:
            print(f"[SKIP] {c}: no usage data")
            continue

        a_bytes = gen_single(
            c, billing_month, csv, used_plist, used_currency,
            rate, bank, phrase, rate_date,
        )
        b_bytes = gen_batch(
            c, billing_month, csv, used_plist, used_currency,
            rate, bank, phrase, rate_date,
        )

        result = compare_excels(a_bytes, b_bytes)

        # 1단계 구조 출력
        struct_ok = (
            result["structure_match"]
            and not result["sheet_diffs"]
        )
        struct_tag = "OK" if struct_ok else "DIFF"
        print(f"[{c}] currency={used_currency}")
        print(f"  Sheets A: {result['a_sheets']}")
        print(f"  Sheets B: {result['b_sheets']}")
        print(f"  Size A: {result['a_info']}")
        print(f"  Size B: {result['b_info']}")
        print(f"  Structure: {struct_tag}")
        if result["sheet_diffs"]:
            for sn, d in result["sheet_diffs"].items():
                print(f"    sheet diff: {sn} → {d}")
                if "missing in A" in d:
                    only_in_b.append((c, sn))
                elif "missing in B" in d:
                    only_in_a.append((c, sn))

        # 2단계 셀 단위 비교
        n_diffs = len(result["cell_diffs"])
        total_cell_diffs += n_diffs
        print(f"  Cell diffs: {n_diffs}")
        if n_diffs > 0:
            for sn, rc, va, vb in result["cell_diffs"][:10]:
                print(f"    [{sn}] cell{rc} A={va!r} B={vb!r}")
            if n_diffs > 10:
                print(f"    ... ({n_diffs - 10} more)")
        print()
        company_results.append((c, n_diffs, struct_ok))

    # 3단계 결과 리포트
    print("=" * 60)
    print(f"=== 최종 결과 ({len(company_results)}개 회사 비교) ===")
    print("=" * 60)
    fully_matched = sum(1 for _, n, s in company_results if n == 0 and s)
    has_diff = [(c, n) for c, n, _ in company_results if n > 0]
    print(f"완전 일치: {fully_matched}개사")
    print(f"불일치 발생: {len(has_diff)}개사")
    print(f"총 불일치 셀 수: {total_cell_diffs}")
    print()
    if only_in_a:
        print("a) 파일 A 에만 있는 시트:")
        for c, sn in only_in_a:
            print(f"   - [{c}] sheet={sn}")
        print()
    if only_in_b:
        print("b) 파일 B 에만 있는 시트:")
        for c, sn in only_in_b:
            print(f"   - [{c}] sheet={sn}")
        print()
    if has_diff:
        print("c) 양쪽에 있지만 값이 다른 셀 — 회사별 요약:")
        for c, n in has_diff:
            print(f"   - [{c}]: {n}개 셀 차이")
        print()
    if total_cell_diffs == 0 and not only_in_a and not only_in_b:
        print(f">>> 전체 {len(company_results)}개 기업, 모든 셀 데이터 100% 일치 <<<")
        return 0
    else:
        print(">>> 불일치 발견 — 위 상세 확인 <<<")
        return 1


if __name__ == "__main__":
    sys.exit(main())
