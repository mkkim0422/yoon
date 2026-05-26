"""validate_invoice_excel 동작 확인.
정상 정산 결과에는 경고 0건, 인위적으로 *0 박은 엑셀에는 경고 발생을 검증.
"""
from __future__ import annotations
import sys, io
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

class _AttrDict(dict):
    def __getattr__(self, k):
        if k in self: return self[k]
        raise AttributeError(k)
    def __setattr__(self, k, v): self[k] = v

class _StubST:
    def __init__(self): self.__dict__["session_state"] = _AttrDict()
    def __getattr__(self, k): return _StubST()
    def __call__(self, *a, **kw): return _StubST()
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def __contains__(self, k): return False
    def __setitem__(self, k, v): pass
    def __getitem__(self, k): return None
    def __iter__(self): return iter([])

sys.modules["streamlit"] = _StubST()
sys.modules["streamlit_sortables"] = _StubST()

from billing.engine import calculate_billing, calculate_billing_by_project
from billing.loader import (
    build_sku_master_from_usage, get_billable_sku_names, load_usage_rows,
    detect_missing_skus,
)
from billing.preprocessor import extract_company_names, preprocess_usage_file
from invoice_generator import generate_formatted_invoice, validate_invoice_excel
from openpyxl import load_workbook

csv = ROOT / "billing.csv"
plist = ROOT / "billing" / "saved_price_list_usd.xlsx"
billing_month = "2026-04"
companies = extract_company_names(str(csv))

# 데이터 있는 첫 USD 회사
company = None
for c in companies:
    try:
        raw = preprocess_usage_file(str(csv), billing_month, company_filter=c)
        if raw:
            company = c; break
    except Exception:
        pass

print(f"테스트 회사: {company}")

raw = preprocess_usage_file(str(csv), billing_month, company_filter=company)
usage = load_usage_rows(raw)
master = build_sku_master_from_usage(usage, str(plist))
billable = get_billable_sku_names(str(plist))
detect_missing_skus(usage, master)
ex = Decimal("1480.80"); mr = Decimal("1.0")
items = calculate_billing(usage, master, ex, mr, mode="account")
proj = calculate_billing_by_project(usage, master, ex, mr, mode="account")
excel_bytes = generate_formatted_invoice(
    line_items=items, company_name=company, billing_month=billing_month,
    exchange_rate=ex, margin_rate=mr, bank_name="하나은행", proj_results=proj,
    price_list_file=str(plist), sku_order=None, currency="USD",
    billable_skus=billable, billing_mode="account",
    per_project_invoices=None, min_charge_amount=0.0,
    min_charge_currency="KRW", rate_date_str="2026.04.30",
    rate_phrase="최종 송금환율 기준", rate_extra="", include_project_sheet=True,
    subtotal_round=2, force_keep_skus=None,
)

# 1) 정상 엑셀 검사
warns = validate_invoice_excel(excel_bytes, line_items=items, company_name=company)
print(f"\n[정상 엑셀] 경고 {len(warns)}건")
for w in warns:
    print(f"  - {w}")
assert len(warns) == 0, f"정상 엑셀에 경고가 발생함: {warns}"
print("  ✓ 정상 엑셀에 경고 없음")

# 2) 일부러 *0 박은 엑셀 검사
wb = load_workbook(io.BytesIO(excel_bytes), data_only=False)
ws = wb["Invoice"]
# Invoice 시트 어딘가에 위험 수식 박기
ws["A100"] = "=ROUND(I77*I78*0,2)"
ws["A101"] = "=SUM()"
ws["A102"] = "=#REF!"
buf = io.BytesIO(); wb.save(buf); bad_bytes = buf.getvalue()
warns = validate_invoice_excel(bad_bytes, line_items=items, company_name=company)
print(f"\n[조작 엑셀] 경고 {len(warns)}건 (기대: >=2건)")
for w in warns:
    print(f"  - {w}")
assert len(warns) >= 2, f"조작 엑셀 경고가 너무 적음: {warns}"
print("  ✓ 조작 엑셀에 경고 발생 확인")

print("\n>>> validator 동작 정상 <<<")
