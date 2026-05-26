"""webapp.py 의 실제 단일/일괄 정산 호출 파라미터를 정확히 재현해 두 엑셀을 비교.

기존 verify_single_vs_batch_excel.py 는 양쪽 모두 margin_rate=0 을 넘겨서
'엑셀 수식에 *0 박히는' 버그를 잡지 못했다. 이 스크립트는:
  - 단일 모드: webapp 라인 2540 의 margin_rate = 1.0 그대로 사용
  - 일괄 모드: webapp 의 _run_batch_single_billing 함수를 ast 로 추출해
    실제 호출부와 동일한 margin_rate=1.0 으로 호출

회사별 결과 엑셀을 셀 단위로 비교 (수식 문자열 포함). 단 한 셀이라도
차이가 있으면 종료코드 1 + 상세 출력.
"""
from __future__ import annotations

import argparse
import io
import sys
from decimal import Decimal
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


# ── Streamlit 무력화 (webapp import 부작용 차단) ──────────────────────
class _AttrDict(dict):
    def __getattr__(self, k):
        if k in self:
            return self[k]
        raise AttributeError(k)
    def __setattr__(self, k, v):
        self[k] = v


class _StubST:
    def __init__(self):
        self.__dict__["session_state"] = _AttrDict()
    def __getattr__(self, k): return _StubST()
    def __call__(self, *a, **kw): return _StubST()
    def __enter__(self): return self
    def __exit__(self, *a): return False
    def __contains__(self, k): return False
    def __setitem__(self, k, v): pass
    def __getitem__(self, k): return None
    def __iter__(self): return iter([])


sys.modules["streamlit"] = _StubST()
sys.modules["streamlit_sortables"] = _StubST()


def _extract_run_batch_single_billing():
    import ast
    src = (ROOT / "webapp.py").read_text(encoding="utf-8")
    tree = ast.parse(src)
    target = next(
        (n for n in tree.body
         if isinstance(n, ast.FunctionDef) and n.name == "_run_batch_single_billing"),
        None,
    )
    assert target, "_run_batch_single_billing not found"
    func_src = ast.get_source_segment(src, target)
    from pathlib import Path as _P
    from decimal import Decimal as _Dec
    from collections import defaultdict as _dd  # noqa: F401
    from billing.engine import calculate_billing as _cb, calculate_billing_by_project as _cbp
    from billing.loader import (
        build_sku_master_from_usage as _bsmf,
        detect_missing_skus as _dms,
        get_free_caps_from_price_list as _gfc,
        load_usage_rows as _lur,
    )
    from billing.models import BillingLineItem as _BLI  # noqa: F401
    from billing.preprocessor import preprocess_usage_file as _puf
    from invoice_generator import generate_formatted_invoice as _gfi
    ns = {
        "Path": _P, "Decimal": _Dec,
        "calculate_billing": _cb, "calculate_billing_by_project": _cbp,
        "build_sku_master_from_usage": _bsmf, "detect_missing_skus": _dms,
        "get_free_caps_from_price_list": _gfc, "load_usage_rows": _lur,
        "preprocess_usage_file": _puf, "generate_formatted_invoice": _gfi,
    }
    exec(func_src, ns)
    return ns["_run_batch_single_billing"]


_run_batch_single_billing = _extract_run_batch_single_billing()

from billing.engine import calculate_billing, calculate_billing_by_project  # noqa: E402
from billing.loader import (  # noqa: E402
    build_sku_master_from_usage, get_billable_sku_names, load_usage_rows,
    detect_missing_skus,
)
from billing.preprocessor import extract_company_names, preprocess_usage_file  # noqa: E402
from invoice_generator import generate_formatted_invoice  # noqa: E402


# ── webapp 실제 단일 정산 호출 흐름 재현 (margin_rate=1.0, line 2540) ──
def gen_single(company, billing_month, csv_path, plist, currency,
               rate, bank, phrase, rate_date):
    raw = preprocess_usage_file(str(csv_path), billing_month, company_filter=company)
    usage = load_usage_rows(raw)
    master = build_sku_master_from_usage(usage, str(plist))
    billable = get_billable_sku_names(str(plist))
    detect_missing_skus(usage, master)  # 동작은 같지만 호출 순서 일치를 위해
    ex = Decimal(str(rate))
    mr = Decimal("1.0")   # ★ webapp 단일 모드 라인 2540 과 동일
    items = calculate_billing(usage, master, ex, mr, mode="account")
    proj = calculate_billing_by_project(usage, master, ex, mr, mode="account")
    return generate_formatted_invoice(
        line_items=items, company_name=company, billing_month=billing_month,
        exchange_rate=ex, margin_rate=mr, bank_name=bank, proj_results=proj,
        price_list_file=str(plist), sku_order=None, currency=currency,
        billable_skus=billable, billing_mode="account",
        per_project_invoices=None, min_charge_amount=0.0,
        min_charge_currency="KRW", rate_date_str=rate_date,
        rate_phrase=phrase, rate_extra="", include_project_sheet=True,
        subtotal_round=0 if currency == "KRW" else 2, force_keep_skus=None,
    )


# ── webapp 실제 일괄 정산 호출 (수정 후 margin_rate=1.0) ──────────────
def gen_batch(company, billing_month, csv_path, plist, currency,
              rate, bank, phrase, rate_date):
    billable = get_billable_sku_names(str(plist))
    res = _run_batch_single_billing(
        selected_company=company, billing_month=billing_month,
        tmp_input_path=Path(csv_path), price_list_file=str(plist),
        currency=currency, exchange_rate=float(rate),
        margin_rate=1.0,    # ★ webapp 일괄 모드 라인 1180/1224/4990/5037 과 동일
        rate_date_str=rate_date, billing_mode="account",
        include_project_sheet=True,
        subtotal_round=0 if currency == "KRW" else 2,
        bank_name=bank, rate_phrase_text=phrase, rate_extra_text="",
        min_charge_amount=0.0, min_charge_currency="KRW",
        sku_order=[], manual_skus=[], hidden_skus=[],
        billable_skus=billable, dl_xlsx=True, dl_pdf=False,
    )
    assert res["ok"], f"batch failed: {res.get('error')}"
    return res["excel_bytes"]


# ── 셀 단위 비교 ────────────────────────────────────────────────────────
def normalize_cell(v):
    if v is None:
        return ""
    if isinstance(v, (int, float)):
        try:
            f = float(v)
            if f == int(f):
                return int(f)
            return round(f, 6)
        except (ValueError, TypeError):
            return v
    return v


def load_excel_structure(xlsx_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    info = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        info[sn] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "cells": {},
            "formulas": {},     # 수식만 별도 카운트
        }
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=r, column=c).value
                info[sn]["cells"][(r, c)] = normalize_cell(v)
                if isinstance(v, str) and v.startswith("="):
                    info[sn]["formulas"][(r, c)] = v
    return info


def compare_excels(a_bytes, b_bytes):
    a = load_excel_structure(a_bytes)
    b = load_excel_structure(b_bytes)
    a_sheets = list(a.keys())
    b_sheets = list(b.keys())
    sheet_diffs = {}
    for sn in set(a_sheets) | set(b_sheets):
        if sn not in a:
            sheet_diffs[sn] = "missing in A"
        elif sn not in b:
            sheet_diffs[sn] = "missing in B"
        elif (a[sn]["max_row"], a[sn]["max_col"]) != (b[sn]["max_row"], b[sn]["max_col"]):
            sheet_diffs[sn] = (
                f"size A={a[sn]['max_row']}x{a[sn]['max_col']} "
                f"B={b[sn]['max_row']}x{b[sn]['max_col']}"
            )
    cell_diffs = []
    formula_diffs = []
    total_cells_a = 0
    total_cells_b = 0
    total_formulas_a = 0
    total_formulas_b = 0
    for sn in a_sheets:
        if sn not in b:
            continue
        total_cells_a += len(a[sn]["cells"])
        total_cells_b += len(b[sn]["cells"])
        total_formulas_a += len(a[sn]["formulas"])
        total_formulas_b += len(b[sn]["formulas"])
        coords = set(a[sn]["cells"].keys()) | set(b[sn]["cells"].keys())
        for rc in sorted(coords):
            va = a[sn]["cells"].get(rc, "")
            vb = b[sn]["cells"].get(rc, "")
            if va != vb:
                cell_diffs.append((sn, rc, va, vb))
                if (isinstance(va, str) and va.startswith("=")) or \
                   (isinstance(vb, str) and vb.startswith("=")):
                    formula_diffs.append((sn, rc, va, vb))
    return {
        "a_sheets": a_sheets, "b_sheets": b_sheets,
        "sheet_diffs": sheet_diffs,
        "cell_diffs": cell_diffs,
        "formula_diffs": formula_diffs,
        "total_cells_a": total_cells_a,
        "total_cells_b": total_cells_b,
        "total_formulas_a": total_formulas_a,
        "total_formulas_b": total_formulas_b,
        "a_info": {sn: (a[sn]["max_row"], a[sn]["max_col"]) for sn in a_sheets},
        "b_info": {sn: (b[sn]["max_row"], b[sn]["max_col"]) for sn in b_sheets},
    }


# ── 메인 ──────────────────────────────────────────────────────────────
def run_round(round_label, targets, billing_month, csv, plist_usd, plist_krw,
              rate, bank, phrase, rate_date, verbose=False):
    print(f"\n{'='*70}")
    print(f"=== {round_label}")
    print(f"{'='*70}")
    total_cell_diffs = 0
    total_formula_diffs = 0
    grand_total_cells = 0
    company_results = []
    for c in targets:
        used_currency = None
        used_plist = None
        for currency, plist in (("USD", plist_usd), ("KRW", plist_krw)):
            try:
                a_bytes = gen_single(
                    c, billing_month, csv, plist, currency,
                    rate, bank, phrase, rate_date,
                )
                wb = openpyxl.load_workbook(io.BytesIO(a_bytes), data_only=False)
                has_data = any(wb[sn].max_row > 5 for sn in wb.sheetnames)
                if has_data:
                    used_currency = currency
                    used_plist = plist
                    break
            except Exception:
                continue
        if not used_currency:
            print(f"[SKIP] {c}: no usage data")
            continue
        a_bytes = gen_single(c, billing_month, csv, used_plist, used_currency,
                              rate, bank, phrase, rate_date)
        b_bytes = gen_batch(c, billing_month, csv, used_plist, used_currency,
                             rate, bank, phrase, rate_date)
        result = compare_excels(a_bytes, b_bytes)
        n_cells = max(result["total_cells_a"], result["total_cells_b"])
        n_form = max(result["total_formulas_a"], result["total_formulas_b"])
        n_diff = len(result["cell_diffs"])
        n_fdiff = len(result["formula_diffs"])
        struct_ok = (result["a_sheets"] == result["b_sheets"] and not result["sheet_diffs"])
        total_cell_diffs += n_diff
        total_formula_diffs += n_fdiff
        grand_total_cells += n_cells
        company_results.append((c, used_currency, n_cells, n_form, n_diff, n_fdiff, struct_ok))
        if verbose or n_diff > 0 or not struct_ok:
            print(f"[{c}] {used_currency} sheets={result['a_sheets']}")
            print(f"  A size={result['a_info']} B size={result['b_info']}")
            print(f"  cells={n_cells} formulas={n_form} cell_diffs={n_diff} formula_diffs={n_fdiff}")
            if result["sheet_diffs"]:
                for sn, d in result["sheet_diffs"].items():
                    print(f"  sheet diff: {sn} -> {d}")
            if n_diff > 0:
                for sn, rc, va, vb in result["cell_diffs"][:15]:
                    print(f"  [{sn}] cell{rc} A={va!r} B={vb!r}")
                if n_diff > 15:
                    print(f"  ... ({n_diff - 15} more)")
    print()
    print(f"--- {round_label} 요약 ---")
    print(f"비교 회사 수: {len(company_results)}")
    print(f"전체 셀 수(합산): {grand_total_cells}")
    print(f"불일치 셀 수: {total_cell_diffs}")
    print(f"불일치 수식 셀 수: {total_formula_diffs}")
    if total_cell_diffs == 0:
        print(f">>> {round_label}: 전체 {grand_total_cells}개 셀 비교 - 일치 <<<")
        return True, grand_total_cells, total_cell_diffs
    else:
        print(f">>> {round_label}: 전체 {grand_total_cells}개 셀 비교 - 불일치 ({total_cell_diffs}개) <<<")
        return False, grand_total_cells, total_cell_diffs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--rounds", type=int, default=1)
    ap.add_argument("--companies", type=int, default=12)
    ap.add_argument("--verbose", action="store_true")
    args = ap.parse_args()

    csv = ROOT / "billing.csv"
    plist_usd = ROOT / "billing" / "saved_price_list_usd.xlsx"
    plist_krw = ROOT / "billing" / "saved_price_list_krw.xlsx"
    assert csv.exists(), f"missing {csv}"

    companies = extract_company_names(str(csv))

    billing_month = "2026-04"
    rate = 1480.80
    bank = "하나은행"
    phrase = "최종 송금환율 기준"
    rate_date = "2026.04.30"

    targets = []
    for c in companies[:60]:
        try:
            raw = preprocess_usage_file(str(csv), billing_month, company_filter=c)
            if raw:
                targets.append(c)
        except Exception:
            pass
        if len(targets) >= args.companies:
            break

    print(f"=== Comparing {len(targets)} companies ===")
    all_pass = True
    summaries = []
    for i in range(1, args.rounds + 1):
        ok, n_cells, n_diffs = run_round(
            f"검증 {i}회차", targets, billing_month, csv, plist_usd, plist_krw,
            rate, bank, phrase, rate_date, verbose=args.verbose,
        )
        summaries.append((i, n_cells, n_diffs, ok))
        if not ok:
            all_pass = False

    print("\n" + "=" * 70)
    print("=== 최종 요약 ===")
    print("=" * 70)
    for i, n_cells, n_diffs, ok in summaries:
        tag = "일치" if ok else f"불일치({n_diffs})"
        print(f"검증 {i}회차: 전체 {n_cells}개 셀 비교 - {tag}")

    return 0 if all_pass else 1


if __name__ == "__main__":
    sys.exit(main())
